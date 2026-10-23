"""Read-only audit of the application's original exports and their saved API sources."""
import json
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

root=Path('output/sample3-validation-20261022')
menu=json.loads((root/'03-menu.json').read_text(encoding='utf-8-sig'))['job']
tc=json.loads((root/'04-test-cases.json').read_text(encoding='utf-8-sig'))['set']
rows=menu['output']['rows'];cases=tc['output']['testCases']
features=[r for r in rows if r.get('included') and r.get('rowType')!='MENU_OBSERVATION']
ids=[t['id'] for t in cases];tc_ids=[t['tcId'] for t in cases]
assert len(ids)==len(set(ids))==len(cases)==tc['output']['caseCount']
assert len(tc_ids)==len(set(tc_ids))
assert {r['rowId'] for r in features}=={t['featureId'] for t in cases if not t['featureId'].startswith('scope:')}
assert all(t['result'] is None for t in cases)
books={}
for filename in ['03-메뉴트리.xlsx','04-TC.xlsx']:
    book=load_workbook(root/filename,read_only=True,data_only=False)
    books[filename]=[{'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'formulaCells':sum(c.data_type=='f' for row in s for c in row)} for s in book]
    assert all(s['formulaCells']==0 for s in books[filename])
    if filename.startswith('04'):
        actual=list(book['TestCase'].values)[1:]
        execution=list(book['실행 정보'].values)[1:]
        assert len(actual)==len(execution)==len(cases)
        assert {row[5] for row in actual}==set(tc_ids)
        assert {row[11] for row in execution}==set(ids)
        assert all(row[12] is None for row in actual)
        depth_mismatch=sum(row[4] in (None,'-') for row,t in zip(actual,cases) if len(t['featurePath'])<5)
source=load_workbook(next(Path('sample/sample3').glob('*.xlsx')),read_only=True,data_only=True)
product='';source_counts=Counter()
for index,row in enumerate(source.active.values,1):
    if index<6:continue
    if row[1]:product=row[1]
    if row[4]:source_counts[product]+=1
summary={'menuStatus':menu['status'],'generationMode':menu['generationMode'],'assistedReview':True,'sourceFunctionCounts':dict(source_counts),
 'menuRows':len(rows),'includedFunctions':len(features),'structureRows':sum(r.get('rowType')=='MENU_OBSERVATION' for r in rows),
 'unassignedFunctions':sum(r.get('rowType')!='MENU_OBSERVATION' and not r.get('included') for r in rows),
 'caseCount':len(cases),'templateCount':dict(Counter(t['featureType'] for t in cases)),
 'executionReadiness':dict(Counter(t.get('executionReadiness') for t in cases)),
 'checks':{'uniqueIds':True,'allIncludedFunctionsHaveTc':True,'tcExcelRowsMatchJson':True,'executionInfoIdsMatch':True,'generatedResultsBlank':True,'originalExportsReadable':True},
 'findings':{'shortPathsWithoutFixedFifthDepth':depth_mismatch,'expectedResultExamples':list(dict.fromkeys(t['expectedResult'] for t in cases))[:5]},'workbooks':books}
(root/'artifact-audit.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False))
